#!/usr/bin/env python3
"""Smoke cached primary-runtime helper paths against the iOS staged shims.

This does not install or register the skills. It extracts the compatibility
packages staged by the iPhone host, places them under a Codex-runtime-shaped
temporary HOME, then runs representative cached helper entrypoints and workbook
APIs against that staged surface.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path


REPO_ROOT = Path(__file__).resolve().parents[1]
RUNTIME_SUPPORT = REPO_ROOT / "Sources/JustBash/OAIPrimaryRuntimeSupport.swift"
PRIMARY_RUNTIME_CACHE = Path("/Users/mweinbach/.codex/plugins/cache/openai-primary-runtime")
PRESENTATIONS_ROOT = (
    PRIMARY_RUNTIME_CACHE
    / "presentations/26.430.10722/skills/presentations"
)
SPREADSHEETS_ROOT = (
    PRIMARY_RUNTIME_CACHE
    / "spreadsheets/26.430.10722/skills/spreadsheets"
)
DOCUMENTS_ROOT = (
    PRIMARY_RUNTIME_CACHE
    / "documents/26.430.10722/skills/documents"
)
IOS_PYTHON_APP = REPO_ROOT / "Apps/JustBashPhone/PythonApp"


SWIFT_STRING_NAMES = [
    "artifactToolPackageJSON",
    "artifactToolCompatModule",
    "presentationJSXCompatModule",
    "presentationJSXRuntimeCompatModule",
    "lucidePackageJSON",
    "lucideCompatModule",
    "sharpPackageJSON",
    "sharpCompatModule",
]


def extract_swift_raw_strings(source: str) -> dict[str, str]:
    values: dict[str, str] = {}
    for name in SWIFT_STRING_NAMES:
        pattern = rf"private static let {re.escape(name)} = #\"\"\"\n(.*?)\n    \"\"\"#"
        match = re.search(pattern, source, flags=re.S)
        if not match:
            raise RuntimeError(f"Could not extract {name} from {RUNTIME_SUPPORT}")
        values[name] = match.group(1)
    return values


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def stage_runtime(home: Path, strings: dict[str, str]) -> Path:
    node_modules = (
        home
        / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules"
    )

    artifact_root = node_modules / "@oai/artifact-tool"
    write_text(artifact_root / "package.json", strings["artifactToolPackageJSON"])
    write_text(artifact_root / "dist/artifact_tool.mjs", strings["artifactToolCompatModule"])
    write_text(artifact_root / "dist/presentation-jsx/index.mjs", strings["presentationJSXCompatModule"])
    write_text(
        artifact_root / "dist/presentation-jsx/jsx-runtime.mjs",
        strings["presentationJSXRuntimeCompatModule"],
    )
    write_text(
        artifact_root / "dist/presentation-jsx/jsx-dev-runtime.mjs",
        strings["presentationJSXRuntimeCompatModule"],
    )

    lucide_root = node_modules / "lucide"
    write_text(lucide_root / "package.json", strings["lucidePackageJSON"])
    write_text(lucide_root / "dist/index.mjs", strings["lucideCompatModule"])

    sharp_root = node_modules / "sharp"
    write_text(sharp_root / "package.json", strings["sharpPackageJSON"])
    write_text(sharp_root / "index.js", strings["sharpCompatModule"])

    return node_modules


def run_command(
    command: list[str],
    *,
    cwd: Path,
    env: dict[str, str],
    label: str,
) -> dict[str, object]:
    proc = subprocess.run(
        command,
        cwd=cwd,
        env=env,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
    )
    return {
        "label": label,
        "command": command,
        "cwd": str(cwd),
        "exitCode": proc.returncode,
        "stdout": proc.stdout,
        "stderr": proc.stderr,
    }


def require_success(result: dict[str, object]) -> None:
    if result["exitCode"] != 0:
        raise RuntimeError(
            "\n".join(
                [
                    f"{result['label']} failed with exit {result['exitCode']}",
                    "command: " + " ".join(str(part) for part in result["command"]),
                    "stdout:",
                    str(result["stdout"]),
                    "stderr:",
                    str(result["stderr"]),
                ]
            )
        )


def smoke_presentation_helper(workdir: Path, env: dict[str, str]) -> dict[str, object]:
    slides_dir = workdir / "slides"
    out_dir = workdir / "presentation-out"
    preview_dir = out_dir / "previews"
    layout_dir = out_dir / "layouts"
    manifest = out_dir / "manifest.json"
    slides_dir.mkdir(parents=True, exist_ok=True)
    write_text(
        slides_dir / "slide-01.mjs",
        """
export async function slide01(presentation, ctx) {
  const slide = presentation.slides.add();
  slide.background.fill = "rgb(255,255,255)";
  ctx.addText(slide, {
    name: "title",
    text: "iOS staged presentation helper smoke",
    left: 48,
    top: 48,
    width: 720,
    height: 96,
    fontSize: 28,
    color: "rgb(17,24,39)",
    fill: "rgb(239,246,255)",
    line: ctx.line("rgb(37,99,235)", 2),
  });
  await ctx.addLucideIcon(slide, {
    icon: "Smartphone",
    left: 48,
    top: 176,
    width: 96,
    height: 96,
    color: "#2563eb",
    name: "phone-icon",
  });
  return slide;
}
""".lstrip(),
    )

    result = run_command(
        [
            "node",
            str(PRESENTATIONS_ROOT / "scripts/build_artifact_deck.mjs"),
            "--slides-dir",
            str(slides_dir),
            "--out",
            str(out_dir / "deck.pptx"),
            "--preview-dir",
            str(preview_dir),
            "--layout-dir",
            str(layout_dir),
            "--manifest",
            str(manifest),
            "--slide-count",
            "1",
            "--scale",
            "0.5",
        ],
        cwd=workdir,
        env=env,
        label="presentations build_artifact_deck",
    )
    require_success(result)

    artifacts = {
        "deck": out_dir / "deck.pptx",
        "preview": preview_dir / "slide-01.png",
        "layout": layout_dir / "slide-01.layout.json",
        "manifest": manifest,
    }
    for name, path in artifacts.items():
        if not path.exists() or path.stat().st_size <= 0:
            raise RuntimeError(f"Presentation smoke missing non-empty {name}: {path}")

    layout = json.loads(artifacts["layout"].read_text(encoding="utf-8"))
    elements = layout.get("elements", [])
    names = [element.get("name") for element in elements]
    has_icon = any(
        element.get("kind") == "image" and "Smartphone" in str(element.get("alt", ""))
        for element in elements
    )
    if "title" not in names or not has_icon:
        raise RuntimeError(f"Presentation layout did not include expected title/icon elements: {elements}")

    return {
        "result": result,
        "artifacts": {key: {"path": str(path), "bytes": path.stat().st_size} for key, path in artifacts.items()},
        "layoutElements": names,
    }


def smoke_lucide_renderer(workdir: Path, env: dict[str, str]) -> dict[str, object]:
    output = workdir / "icons" / "smartphone.png"
    result = run_command(
        [
            "node",
            str(PRESENTATIONS_ROOT / "scripts/render_lucide_icon.mjs"),
            "--icon",
            "Smartphone",
            "--output",
            str(output),
            "--size",
            "64",
        ],
        cwd=workdir,
        env=env,
        label="presentations render_lucide_icon",
    )
    require_success(result)
    if not output.exists() or output.stat().st_size <= 50:
        raise RuntimeError(f"Lucide PNG smoke did not produce a useful PNG: {output}")
    return {"result": result, "artifact": {"path": str(output), "bytes": output.stat().st_size}}


def smoke_spreadsheet_api(workdir: Path, env: dict[str, str]) -> dict[str, object]:
    script = workdir / "spreadsheet-smoke.mjs"
    out_dir = workdir / "spreadsheet-out"
    write_text(
        script,
        """
import fs from "node:fs/promises";
import { SpreadsheetFile, Workbook } from "@oai/artifact-tool";

const outDir = new URL("./spreadsheet-out/", import.meta.url);
await fs.mkdir(outDir, { recursive: true });

const workbook = Workbook.create();
const sheet = workbook.worksheets.add("Smoke");
sheet.getRange("A1:B3").values = [
  ["label", "value"],
  ["alpha", 2],
  ["beta", 3],
];
sheet.getRange("D1:D3").formulas = [["=SUM(B2:B3)"], ["=AVERAGE(B2:B3)"], ["=D1*2"]];
if (sheet.getRange("D1:D3").values[0][0] !== 5) throw new Error("SUM formula failed");
if (sheet.getRange("D1:D3").values[2][0] !== 10) throw new Error("dependent formula failed");
const errors = await workbook.inspect({ kind: "formula" });
if (errors.errors.length) throw new Error("formula inspection found errors");
if (!JSON.parse(workbook.trace("Smoke!D3").ndjson).dependencies.length) {
  throw new Error("trace did not include dependencies");
}
const chart = sheet.charts.add("line", sheet.getRange("A1:B3"));
chart.setPosition("F1", "L12");
chart.title = "Smoke Chart";
sheet.tables.add("A1:B3", true, "SmokeTable");
const preview = await workbook.render({ sheetName: "Smoke", range: "A1:L12", scale: 1 });
await preview.save(new URL("preview.png", outDir).pathname);
const xlsx = await SpreadsheetFile.exportXlsx(workbook);
await xlsx.save(new URL("workbook.xlsx", outDir).pathname);
const imported = await SpreadsheetFile.importXlsx(xlsx);
if (imported.worksheets.getItem("Smoke").getRange("A1").values[0][0] !== "label") {
  throw new Error("round-trip import failed");
}
const xlsxText = Buffer.from(await xlsx.arrayBuffer()).toString("latin1");
if (!xlsxText.includes("xl/charts/chart1.xml")) throw new Error("chart XML missing");
console.log(JSON.stringify({ ok: true }));
""".lstrip(),
    )
    result = run_command(
        ["node", str(script)],
        cwd=workdir,
        env=env,
        label="spreadsheets artifact-tool api",
    )
    require_success(result)

    artifacts = {
        "xlsx": out_dir / "workbook.xlsx",
        "preview": out_dir / "preview.png",
    }
    for name, path in artifacts.items():
        if not path.exists() or path.stat().st_size <= 0:
            raise RuntimeError(f"Spreadsheet smoke missing non-empty {name}: {path}")

    return {
        "result": result,
        "artifacts": {key: {"path": str(path), "bytes": path.stat().st_size} for key, path in artifacts.items()},
    }


def write_minimal_docx(path: Path) -> None:
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>
</Types>
"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>
"""
    document_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdComments" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>
</Relationships>
"""
    document = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p>
      <w:commentRangeStart w:id="0"/>
      <w:r><w:t>Hello</w:t></w:r>
      <w:commentRangeEnd w:id="0"/>
      <w:r><w:commentReference w:id="0"/></w:r>
    </w:p>
  </w:body>
</w:document>
"""
    comments = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:comment w:id="0" w:author="JustBash"><w:p><w:r><w:t>Remove me</w:t></w:r></w:p></w:comment>
</w:comments>
"""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/_rels/document.xml.rels", document_rels)
        zf.writestr("word/document.xml", document)
        zf.writestr("word/comments.xml", comments)


def write_fake_soffice(path: Path) -> None:
    write_text(
        path,
        """
#!/usr/bin/env python3
from pathlib import Path
import sys

args = sys.argv[1:]
try:
    outdir = Path(args[args.index("--outdir") + 1])
    convert_to = args[args.index("--convert-to") + 1].split(":", 1)[0].lower()
    input_path = Path(args[-1])
except Exception as exc:
    print(f"unsupported soffice command: {exc}", file=sys.stderr)
    raise SystemExit(1)

outdir.mkdir(parents=True, exist_ok=True)
out = outdir / f"{input_path.stem}.{convert_to}"
if convert_to == "pdf":
    out.write_text(
        "%PDF-1.4\\n"
        "% JUSTBASH_PAGE_SIZE 612 x 792 pts\\n"
        "1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\\n"
        "2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\\n"
        "3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << >> /Contents 4 0 R >> endobj\\n"
        "4 0 obj << /Length 0 >> stream\\n\\nendstream endobj\\n"
        "trailer << /Root 1 0 R >>\\n%%EOF\\n",
        encoding="latin1",
    )
else:
    out.write_bytes(b"PK\\x03\\x04JustBash compatibility placeholder\\n")
print(f"convert {input_path} -> {out}")
""".lstrip(),
    )
    path.chmod(0o755)


def smoke_documents_lxml_helpers(workdir: Path, env: dict[str, str]) -> dict[str, object]:
    input_docx = workdir / "documents-in.docx"
    protected_docx = workdir / "documents-protected.docx"
    stripped_docx = workdir / "documents-stripped.docx"
    xlsx_path = workdir / "documents-table.xlsx"
    table_docx = workdir / "documents-table.docx"
    patched_table_docx = workdir / "documents-table-patched.docx"
    table_csv = workdir / "documents-table.csv"
    render_dir = workdir / "documents-render"
    fake_bin = workdir / "bin"
    fake_bin.mkdir(parents=True, exist_ok=True)
    write_fake_soffice(fake_bin / "soffice")
    write_minimal_docx(input_docx)

    python_env = env.copy()
    existing_path = python_env.get("PYTHONPATH")
    python_paths = [str(IOS_PYTHON_APP), str(IOS_PYTHON_APP / "site-packages")]
    if existing_path:
        python_paths.append(existing_path)
    python_env["PYTHONPATH"] = os.pathsep.join(python_paths)
    python_env["PATH"] = str(fake_bin) + os.pathsep + python_env.get("PATH", "")

    protection = run_command(
        [
            sys.executable,
            str(DOCUMENTS_ROOT / "scripts/set_protection.py"),
            str(input_docx),
            "--mode",
            "readOnly",
            "--out",
            str(protected_docx),
        ],
        cwd=workdir,
        env=python_env,
        label="documents set_protection",
    )
    require_success(protection)

    strip = run_command(
        [
            sys.executable,
            str(DOCUMENTS_ROOT / "scripts/comments_strip.py"),
            str(protected_docx),
            "--out",
            str(stripped_docx),
        ],
        cwd=workdir,
        env=python_env,
        label="documents comments_strip",
    )
    require_success(strip)

    with zipfile.ZipFile(stripped_docx) as zf:
        names = set(zf.namelist())
        settings = zf.read("word/settings.xml").decode("utf-8")
        document = zf.read("word/document.xml").decode("utf-8")
        content_types = zf.read("[Content_Types].xml").decode("utf-8")
        rels = zf.read("word/_rels/document.xml.rels").decode("utf-8")

    if "word/comments.xml" in names:
        raise RuntimeError("comments_strip did not remove word/comments.xml")
    for marker in ["commentRangeStart", "commentRangeEnd", "commentReference"]:
        if marker in document:
            raise RuntimeError(f"comments_strip left {marker} in document.xml")
    if "comments" in rels or "comments.xml" in content_types:
        raise RuntimeError("comments_strip left comment relationships/content types")
    if "documentProtection" not in settings or "readOnly" not in settings:
        raise RuntimeError("set_protection did not create readOnly settings.xml")

    write_text(
        workdir / "make-table-xlsx.py",
        """
from openpyxl import Workbook
from pathlib import Path
import sys

out = Path(sys.argv[1])
wb = Workbook()
ws = wb.active
ws.title = "Data"
ws.append(["name", "value"])
ws.append(["alpha", 2])
ws.append(["beta", 3])
wb.save(out)
""".lstrip(),
    )
    make_xlsx = run_command(
        [sys.executable, str(workdir / "make-table-xlsx.py"), str(xlsx_path)],
        cwd=workdir,
        env=python_env,
        label="documents make xlsx fixture",
    )
    require_success(make_xlsx)
    xlsx_to_docx = run_command(
        [
            sys.executable,
            str(DOCUMENTS_ROOT / "scripts/xlsx_to_docx_table.py"),
            str(xlsx_path),
            "--out",
            str(table_docx),
            "--title",
            "Smoke Table",
        ],
        cwd=workdir,
        env=python_env,
        label="documents xlsx_to_docx_table",
    )
    require_success(xlsx_to_docx)
    docx_to_csv = run_command(
        [
            sys.executable,
            str(DOCUMENTS_ROOT / "scripts/docx_table_to_csv.py"),
            str(table_docx),
            "--out",
            str(table_csv),
        ],
        cwd=workdir,
        env=python_env,
        label="documents docx_table_to_csv",
    )
    require_success(docx_to_csv)
    csv_text = table_csv.read_text(encoding="utf-8").replace("\r\n", "\n")
    if csv_text.strip() != "name,value\nalpha,2\nbeta,3":
        raise RuntimeError(f"table DOCX round-trip CSV did not match expected content: {csv_text!r}")
    patch_docx = run_command(
        [
            sys.executable,
            str(DOCUMENTS_ROOT / "scripts/docx_ooxml_patch.py"),
            str(table_docx),
            "--out",
            str(patched_table_docx),
            "--header-date",
            "May 1, 2026",
            "--add-page-numbers",
            "--hyperlink-first",
            "https://example.com",
        ],
        cwd=workdir,
        env=python_env,
        label="documents docx_ooxml_patch python-docx path",
    )
    require_success(patch_docx)
    with zipfile.ZipFile(patched_table_docx) as zf:
        patched_names = set(zf.namelist())
        patched_doc = zf.read("word/document.xml").decode("utf-8")
        patched_rels = zf.read("word/_rels/document.xml.rels").decode("utf-8")
    if not {"word/header1.xml", "word/footer1.xml"}.issubset(patched_names):
        raise RuntimeError("docx_ooxml_patch did not materialize header/footer parts")
    if "w:hyperlink" not in patched_doc or "https://example.com" not in patched_rels:
        raise RuntimeError("docx_ooxml_patch did not create the expected external hyperlink")

    render_docx = run_command(
        [
            sys.executable,
            str(DOCUMENTS_ROOT / "render_docx.py"),
            str(patched_table_docx),
            "--output_dir",
            str(render_dir),
            "--width",
            "800",
            "--height",
            "1000",
        ],
        cwd=workdir,
        env=python_env,
        label="documents render_docx bounded iOS path",
    )
    require_success(render_docx)
    rendered_page = render_dir / "page-1.png"
    if not rendered_page.exists() or rendered_page.stat().st_size <= 50:
        raise RuntimeError(f"render_docx did not create a useful page PNG: {rendered_page}")
    if rendered_page.read_bytes()[:8] != b"\x89PNG\r\n\x1a\n":
        raise RuntimeError("render_docx output is not a PNG")

    return {
        "result": protection,
        "secondResult": strip,
        "thirdResult": xlsx_to_docx,
        "fourthResult": docx_to_csv,
        "fifthResult": patch_docx,
        "sixthResult": render_docx,
        "artifacts": {
            "protected": {"path": str(protected_docx), "bytes": protected_docx.stat().st_size},
            "stripped": {"path": str(stripped_docx), "bytes": stripped_docx.stat().st_size},
            "tableDocx": {"path": str(table_docx), "bytes": table_docx.stat().st_size},
            "patchedTableDocx": {"path": str(patched_table_docx), "bytes": patched_table_docx.stat().st_size},
            "tableCsv": {"path": str(table_csv), "bytes": table_csv.stat().st_size},
            "renderedPage": {"path": str(rendered_page), "bytes": rendered_page.stat().st_size},
        },
    }


def build_report(workdir: Path) -> dict[str, object]:
    if not RUNTIME_SUPPORT.exists():
        raise FileNotFoundError(RUNTIME_SUPPORT)
    if not PRESENTATIONS_ROOT.exists():
        raise FileNotFoundError(PRESENTATIONS_ROOT)
    if not SPREADSHEETS_ROOT.exists():
        raise FileNotFoundError(SPREADSHEETS_ROOT)
    if not DOCUMENTS_ROOT.exists():
        raise FileNotFoundError(DOCUMENTS_ROOT)

    strings = extract_swift_raw_strings(RUNTIME_SUPPORT.read_text(encoding="utf-8"))
    home = workdir / "home"
    node_modules = stage_runtime(home, strings)
    env = os.environ.copy()
    env["HOME"] = str(home)
    env["PPTX_COMEBACK_REQUIRE_ROOTS"] = str(node_modules)

    checks: list[dict[str, object]] = []
    for label, fn in [
        ("documents.ooxml_docx_helpers", smoke_documents_lxml_helpers),
        ("presentations.build_artifact_deck", smoke_presentation_helper),
        ("presentations.render_lucide_icon", smoke_lucide_renderer),
        ("spreadsheets.artifact_tool_api", smoke_spreadsheet_api),
    ]:
        try:
            detail = fn(workdir, env)
            checks.append({"name": label, "status": "ok", **detail})
        except Exception as exc:  # noqa: BLE001 - diagnostic script.
            checks.append({"name": label, "status": "failed", "error": str(exc)})

    return {
        "overall": "ok" if all(check["status"] == "ok" for check in checks) else "failed",
        "workdir": str(workdir),
        "stagedNodeModules": str(node_modules),
        "checks": checks,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--json", action="store_true", help="emit JSON report")
    parser.add_argument(
        "--keep-workdir",
        action="store_true",
        help="keep temporary artifacts after the smoke run",
    )
    args = parser.parse_args()

    with tempfile.TemporaryDirectory(prefix="justbash-primary-runtime-smoke-") as td:
        workdir = Path(td)
        report = build_report(workdir)
        if args.keep_workdir:
            kept = Path(tempfile.mkdtemp(prefix="justbash-primary-runtime-smoke-kept-"))
            subprocess.run(["cp", "-R", f"{workdir}/.", str(kept)], check=True)
            report["keptWorkdir"] = str(kept)

    if args.json:
        print(json.dumps(report, indent=2, sort_keys=True))
    else:
        print(f"Primary runtime helper smoke: {report['overall']}")
        for check in report["checks"]:
            print(f"- {check['name']}: {check['status']}")
            if check["status"] != "ok":
                print(f"  {check.get('error', '')}")
    return 0 if report["overall"] == "ok" else 1


if __name__ == "__main__":
    raise SystemExit(main())
